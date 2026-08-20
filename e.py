import os
import time
import json
import threading
import requests
import re
import random
import sqlite3
import openpyxl
from collections import Counter
from tg_helper import (
    send_telegram_message, delete_webhook, get_telegram_updates, 
    escape_html, edit_telegram_message, answer_callback_query, 
    send_tg_reply, send_tg_document
)

# ==============================================================================
# KONFIGURASI BOT & API
# ==============================================================================
TELEGRAM_BOT_TOKEN = "8833933759:AAGyGo2CqskLPBK1KuFyUluWxsbKgajVTGU"
CHAT_ID = "1003914110525"

# DAFTAR ID TELEGRAM ADMIN
ADMIN_IDS = [8525945799]

GROUP_CHAT_URL = "https://t.me/link_grup_kamu"    # Link tombol CHAT ↗
CHANNEL_URL = "https://t.me/link_channel_kamu"   # Link tombol NUMBER ↗

# API KEY & URL
API_1_URL = "https://www.flashsms.space/api/cdr/viewstats?api_token=-W2DGJqcAhOs9m-crCa9wGZ9bOtqmhgzrZHSoHcczHY"
API_2_URL = "https://www.ksiiprn.com/api/v1/iprn/messages"
API_2_TOKEN = "sk_live_QDSlkOM5NIFEXSqxFFKIqWv3anBEZtA53eiIsjal"
API_3_URL = "https://augestel.com/api/v1/iprn/messages"
API_3_TOKEN = "sk_live_Yeg9JqxCUiYR7wSJ9bjShzTZwF7f4qMWSnn3ivaj"

# API 4 (THIRDWAVE) - Endpoint Traffic & Token Baru
API_4_URL = "https://app.thirdwave.im/api/v1/traffic"
API_4_TOKEN = "tw_live_6258a85c82d34fd99f4ac343418e0bdf688a6ffa96c8cd135d5012f05599a5bd"

CHECK_INTERVAL = 4
KSIIPRN_CHECK_INTERVAL = 30
AUGESTEL_CHECK_INTERVAL = 15
THIRDWAVE_CHECK_INTERVAL = 10
TOP_RCV_INTERVAL = 7200

cache = set()
last_update_id = 0
country_counter = Counter()
last_top_rcv_time = time.time()

ksiiprn_cooldown_until = 0
last_ksiiprn_check = 0
ksiiprn_backoff_factor = 1

augestel_cooldown_until = 0
last_augestel_check = 0

thirdwave_cooldown_until = 0

# HTTP Session Global
http_session = requests.Session()

# ==============================================================================
# DATABASE SETUP (SQLite)
# ==============================================================================
DB_PATH = "bot_otp.db"

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS numbers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            country TEXT NOT NULL,
            phone_number TEXT UNIQUE NOT NULL
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        )
    ''')
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('otp_channel', 'https://t.me/link_channel_kamu')")
    conn.commit()
    conn.close()

init_db()

def get_setting(key, default=""):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
    row = cursor.fetchone()
    conn.close()
    return row[0] if row else default

def set_setting(key, value):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, value))
    conn.commit()
    conn.close()

# ==============================================================================
# BENDERA & COUNTRY MAPPING
# ==============================================================================
def load_flags():
    if os.path.exists("flag.json"):
        try:
            with open("flag.json", "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as e:
            print("⚠️ Gagal membaca flag.json:", e)
    return {}

FLAG_DATA = load_flags()

def get_country_info(number):
    if not number or number == "-":
        return {"name": "UNKNOWN", "flag": "🌐", "code": "INT"}
    
    num_str = str(number).lstrip("+")
    sorted_prefixes = sorted(FLAG_DATA.keys(), key=len, reverse=True)
    for prefix in sorted_prefixes:
        if num_str.startswith(prefix):
            return FLAG_DATA[prefix]
            
    return {"name": "INTERNATIONAL", "flag": "🌐", "code": "INT"}

def get_flag_by_country_name(country_name):
    c_upper = str(country_name).upper().strip()
    for prefix, data in FLAG_DATA.items():
        if data.get("name", "").upper() == c_upper or data.get("code", "").upper() == c_upper:
            return data.get("flag", "🌐")
    return "🌐"

def extract_otp(text):
    if not text:
        return "-"
    hyphen_match = re.search(r'\b\d{3,4}[- ]\d{3,4}\b', str(text))
    if hyphen_match:
        return hyphen_match.group(0).replace('-', '').replace(' ', '')
    match = re.search(r'\b\d{4,8}\b', str(text))
    if match:
        return match.group(0)
    return "-"

def mask_number_crown(number):
    if not number or number == "-":
        return "-"
    str_num = str(number).strip().lstrip("+")
    if len(str_num) >= 8:
        return f"{str_num[:4]}👑{str_num[-4:]}"
    elif len(str_num) > 4:
        mid = len(str_num) // 2
        return f"{str_num[:mid]}👑{str_num[mid:]}"
    return str_num

# ==============================================================================
# KEYBOARD & UI MENU GENERATOR
# ==============================================================================
def get_main_reply_keyboard():
    return {
        "keyboard": [
            [{"text": "💥 Get Number"}, {"text": "🔥 Get File"}],
            [{"text": "🏫 Chat Support"}, {"text": "📊 Top RCV"}]
        ],
        "resize_keyboard": True
    }

def get_whatsapp_stock_view():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM numbers")
    total_stock = cursor.fetchone()[0]
    conn.close()

    text = (
        f"📱 <b>WHATSAPP OTP</b>\n\n"
        f"• <b>Stok Tersedia:</b> <code>{total_stock} Nomor</code>\n"
        f"• <b>Status Server:</b> 🟢 <code>Active</code>"
    )
    keyboard = {
        "inline_keyboard": [
            [{"text": f"🟢 Ambil Nomor ({total_stock})", "callback_data": "show_countries"}]
        ]
    }
    return text, keyboard

def get_countries_view():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("SELECT country, COUNT(*) FROM numbers GROUP BY country")
    rows = cursor.fetchall()
    conn.close()

    if not rows:
        return "⚠️ Belum ada stok nomor yang tersedia.", None

    inline_keyboard = []
    for country, count in rows:
        flag = get_flag_by_country_name(country)
        btn_text = f"{flag} {country.upper()} ({count})"
        inline_keyboard.append([{"text": btn_text, "callback_data": f"getnum_{country}"}])

    return "Pilih negara yang ingin Anda ambil nomornya:", {"inline_keyboard": inline_keyboard}

def get_numbers_view(country):
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute("SELECT phone_number FROM numbers WHERE UPPER(country) = ?", (country.upper(),))
    rows = cursor.fetchall()
    conn.close()

    if not rows:
        return f"❌ Stok nomor untuk negara <b>{country.upper()}</b> habis.", None

    all_numbers = [r[0] for r in rows]
    selected_numbers = random.sample(all_numbers, min(5, len(all_numbers)))

    flag = get_flag_by_country_name(country)
    otp_channel = get_setting("otp_channel", "https://t.me/link_channel_kamu")

    text = (
        f"💬 {flag} <b>{country.upper()} 5 Number:</b>\n"
        f"<i>Menunggu OTP masuk... (Sesi: 30 Menit)</i>\n"
        f"<i>Last sync: {time.strftime('%H:%M:%S')}</i>"
    )

    inline_keyboard = []
    for num in selected_numbers:
        inline_keyboard.append([{"text": f"📱 {num}", "copy_text": {"text": num}}])

    inline_keyboard.append([
        {"text": "🔄 Change Number", "callback_data": f"getnum_{country}"},
        {"text": "🌐 Change Country", "callback_data": "show_countries"}
    ])
    inline_keyboard.append([
        {"text": "✉️ Link Group OTP", "url": otp_channel}
    ])

    return text, {"inline_keyboard": inline_keyboard}

# ==============================================================================
# MONITORING & PROCESS SMS
# ==============================================================================
def process_messages(messages, source_name="API"):
    for item in messages:
        if not isinstance(item, dict):
            continue
        
        date_time = item.get("date") or item.get("created_at") or item.get("timestamp") or item.get("time") or ""
        number = (
            item.get("number") or item.get("phone") or item.get("to") or 
            item.get("destination") or item.get("msisdn") or item.get("cli") or "-"
        )
        msg_text = (
            item.get("message") or item.get("msg") or item.get("text") or 
            item.get("sms") or item.get("content") or item.get("body") or 
            item.get("last_message") or ""
        )
        
        msg_id = str(item.get("id") or (str(date_time) + str(number) + str(msg_text[:10])))
        
        if not msg_id or msg_id in cache:
            continue
        cache.add(msg_id)

        if len(cache) > 3000:
            cache.pop()

        if not msg_text and number == "-":
            continue

        otp = extract_otp(msg_text)
        country = get_country_info(number)
        masked_num = mask_number_crown(number)

        message = f"{country['flag']} {country['code']} 💬 {masked_num} #EN"
        send_telegram_message(
            TELEGRAM_BOT_TOKEN, CHAT_ID, message, otp=otp,
            group_url=GROUP_CHAT_URL, channel_url=CHANNEL_URL
        )

        country_name = str(country.get('name', 'UNKNOWN')).upper()
        country_key = f"{country['flag']} {country_name}"
        country_counter[country_key] += 1

        print(f"[{time.strftime('%H:%M:%S')}] 🔥 [{source_name}] {country['flag']} {country['code']} {masked_num} | OTP: {otp}")

def fetch_api_1():
    try:
        headers = {"Content-Type": "application/json"}
        res = http_session.get(API_1_URL, headers=headers, timeout=25)
        data = res.json()
        messages = []
        if isinstance(data, dict):
            messages = data.get("results") or next((v for v in data.values() if isinstance(v, list)), [])
        elif isinstance(data, list):
            messages = data
        process_messages(messages, source_name="FlashSMS")
    except requests.exceptions.Timeout:
        print(f"[{time.strftime('%H:%M:%S')}] ⏳ API 1 (FlashSMS) Timeout (>25s)")
    except Exception as e:
        print(f"[{time.strftime('%H:%M:%S')}] ⚠️ API 1 (FlashSMS) Error:", e)

def fetch_api_2():
    global ksiiprn_cooldown_until, ksiiprn_backoff_factor
    if time.time() < ksiiprn_cooldown_until:
        return

    try:
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {API_2_TOKEN}"
        }
        res = http_session.get(API_2_URL, headers=headers, timeout=25)
        
        if res.status_code == 429:
            retry_after = res.headers.get("Retry-After") or res.headers.get("retry-after")
            cooldown_time = int(retry_after) if retry_after and retry_after.isdigit() else min(60 * ksiiprn_backoff_factor, 600)
            ksiiprn_backoff_factor *= 2
            ksiiprn_cooldown_until = time.time() + cooldown_time
            print(f"[{time.strftime('%H:%M:%S')}] ⏳ KSIIPRN Rate Limit 429! Cooldown {cooldown_time}s...")
            return

        if res.status_code != 200:
            return

        ksiiprn_backoff_factor = 1
        data = res.json()
        messages = data.get("data") or data.get("messages") or data.get("results") or [] if isinstance(data, dict) else data
        if messages:
            process_messages(messages, source_name="KSIIPRN")

    except requests.exceptions.Timeout:
        print(f"[{time.strftime('%H:%M:%S')}] ⏳ API 2 (KSIIPRN) Timeout (>25s)")
    except Exception as e:
        print(f"[{time.strftime('%H:%M:%S')}] ⚠️ API 2 (KSIIPRN) Error:", e)

def fetch_api_3():
    global augestel_cooldown_until
    if time.time() < augestel_cooldown_until:
        return

    try:
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {API_3_TOKEN}"
        }
        res = http_session.get(API_3_URL, headers=headers, timeout=25)
        
        if res.status_code == 429:
            augestel_cooldown_until = time.time() + 60
            print(f"[{time.strftime('%H:%M:%S')}] ⏳ AUGESTEL Rate Limit 429! Cooldown 60s...")
            return

        if res.status_code != 200:
            return

        data = res.json()
        messages = data.get("data") or data.get("messages") or data.get("results") or [] if isinstance(data, dict) else data
        if messages:
            process_messages(messages, source_name="AUGESTEL")

    except requests.exceptions.Timeout:
        print(f"[{time.strftime('%H:%M:%S')}] ⏳ API 3 (AUGESTEL) Timeout (>25s)")
    except Exception as e:
        print(f"[{time.strftime('%H:%M:%S')}] ⚠️ API 3 (AUGESTEL) Error:", e)

def fetch_api_4():
    global thirdwave_cooldown_until
    if time.time() < thirdwave_cooldown_until:
        return

    try:
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {API_4_TOKEN}"
        }
        res = http_session.get(API_4_URL, headers=headers, timeout=25)
        
        if res.status_code == 429:
            thirdwave_cooldown_until = time.time() + 60
            print(f"[{time.strftime('%H:%M:%S')}] ⏳ THIRDWAVE Rate Limit 429! Cooldown 60s...")
            return

        if res.status_code != 200:
            print(f"[{time.strftime('%H:%M:%S')}] ⚠️ API 4 (THIRDWAVE) Status Code: {res.status_code}")
            return

        data = res.json()
        rows = data.get("rows", [])
        
        # Mapping khusus Thirdwave Endpoint /traffic
        messages = []
        for item in rows:
            if isinstance(item, dict):
                messages.append({
                    "id": item.get("id"),
                    "number": item.get("destinationNumber"),
                    "message": item.get("messageBody"),
                    "date": item.get("receivedAt")
                })

        if messages:
            process_messages(messages, source_name="Thirdwave")

    except requests.exceptions.Timeout:
        print(f"[{time.strftime('%H:%M:%S')}] ⏳ API 4 (THIRDWAVE) Timeout (>25s)")
    except Exception as e:
        print(f"[{time.strftime('%H:%M:%S')}] ⚠️ API 4 (THIRDWAVE) Error:", e)

# ==============================================================================
# LAPORAN LIVE TRAFFIC
# ==============================================================================
def generate_top_rcv_text():
    global country_counter
    if not country_counter:
        return (
            "👑 <b>LIVE TRAFFIC ANALYTICS</b>\n"
            "───────────────────────────\n"
            "💬 <i>Belum ada traffic SMS pada sesi ini.</i>\n"
            "───────────────────────────\n"
            "⏳ <i>Traffic direset otomatis setiap 2 jam</i>"
        )
    
    top_countries = country_counter.most_common(10)
    total_sms = sum(country_counter.values())
    lines = []
    medals = ["🥇", "🥈", "🥉"]
    
    for idx, (country_info, count) in enumerate(top_countries, 1):
        percentage = (count / total_sms) * 100
        badge = medals[idx - 1] if idx <= 3 else "▫️"
        lines.append(f"{badge} <b>{country_info}</b> ── <code>{count} SMS</code> (<b>{percentage:.1f}%</b>)")
    
    list_str = "\n".join(lines)
    
    return (
        f"👑 <b>LIVE TRAFFIC ANALYTICS</b>\n"
        f"───────────────────────────\n"
        f"{list_str}\n"
        f"───────────────────────────\n"
        f"📊 <b>Total Traffic:</b> <code>{total_sms} SMS</code>\n"
        f"⏳ <i>Traffic direset otomatis setiap 2 jam</i>"
    )

def send_top_rcv_report(target_chat=CHAT_ID, reset=True):
    global country_counter
    caption = generate_top_rcv_text()
    send_telegram_message(
        TELEGRAM_BOT_TOKEN, target_chat, caption, otp=None, 
        group_url=GROUP_CHAT_URL, channel_url=CHANNEL_URL
    )
    if reset:
        country_counter.clear()

# ==============================================================================
# HELPER ADMIN & CONVERTER FILE
# ==============================================================================
def handle_file_conversion(chat_id, file_id, file_name):
    try:
        get_file_url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/getFile?file_id={file_id}"
        res = http_session.get(get_file_url, timeout=10).json()
        
        if not res.get("ok"):
            send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "❌ Gagal mengunduh file dari Telegram.")
            return

        file_path_tg = res["result"]["file_path"]
        download_url = f"https://api.telegram.org/file/bot{TELEGRAM_BOT_TOKEN}/{file_path_tg}"

        local_excel = f"temp_{int(time.time())}_{file_name}"
        excel_data = http_session.get(download_url, timeout=30).content
        with open(local_excel, "wb") as f:
            f.write(excel_data)

        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "⏳ <b>Sedang mengekstrak dan memilah nomor HP per negara...</b>")

        wb = openpyxl.load_workbook(local_excel, data_only=True)
        extracted_numbers = []
        phone_keywords = ['number', 'phone', 'msisdn', 'nomor', 'no', 'cli', 'destination', 'to', 'mobile', 'telp', 'nohp', 'no_hp']

        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
            target_indices = [idx for idx, col_name in enumerate(header) if any(kw in col_name for kw in phone_keywords)]

            if target_indices:
                for row in rows[1:]:
                    for idx in target_indices:
                        if idx < len(row) and row[idx] is not None:
                            val = str(row[idx]).strip()
                            if val.endswith('.0'): 
                                val = val[:-2]
                            cleaned = re.sub(r'[^\d+]', '', val)
                            if cleaned:
                                extracted_numbers.append(cleaned)
            else:
                for row in rows:
                    for cell in row:
                        if cell is not None:
                            val = str(cell).strip()
                            if val.endswith('.0'): 
                                val = val[:-2]
                            cleaned = re.sub(r'[^\d+]', '', val)
                            if 7 <= len(cleaned.lstrip('+')) <= 15 and cleaned.lstrip('+').isdigit():
                                extracted_numbers.append(cleaned)

        if not extracted_numbers:
            send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "⚠️ <b>Tidak ada nomor HP yang ditemukan di dalam file Excel tersebut.</b>")
            if os.path.exists(local_excel):
                os.remove(local_excel)
            return

        country_groups = {}
        for num in extracted_numbers:
            info = get_country_info(num)
            c_name = info.get("name", "UNKNOWN").title()
            c_flag = info.get("flag", "🌐")

            if c_name not in country_groups:
                country_groups[c_name] = {"flag": c_flag, "numbers": []}
            
            country_groups[c_name]["numbers"].append(num)

        for c_name, data in country_groups.items():
            clean_c_name = re.sub(r'[^\w\s-]', '', c_name).strip().replace(" ", "_")
            txt_filename = f"{clean_c_name}_Numbers.txt"
            local_txt = f"temp_{int(time.time())}_{txt_filename}"

            with open(local_txt, "w", encoding="utf-8") as txt_file:
                txt_file.write("\n".join(data["numbers"]) + "\n")

            caption = f"{data['flag']} Region: {c_name}\n📊 Total Count: {len(data['numbers'])}"
            send_tg_document(TELEGRAM_BOT_TOKEN, chat_id, local_txt, caption=caption)

            if os.path.exists(local_txt):
                os.remove(local_txt)

        if os.path.exists(local_excel):
            os.remove(local_excel)

    except Exception as e:
        print("⚠️ Error saat ekstrak nomor file:", e)
        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, f"❌ <b>Terjadi kesalahan:</b>\n<code>{escape_html(str(e))}</code>")

# ==============================================================================
# TELEGRAM BOT LISTENER
# ==============================================================================
def telegram_listener():
    global last_update_id
    while True:
        try:
            updates = get_telegram_updates(TELEGRAM_BOT_TOKEN, last_update_id + 1)
            for update in updates:
                last_update_id = update["update_id"]

                # 1. CALLBACK QUERY HANDLER
                if "callback_query" in update:
                    cb = update["callback_query"]
                    cb_id = cb["id"]
                    chat_id = cb["message"]["chat"]["id"]
                    chat_type = cb["message"]["chat"].get("type", "")
                    msg_id = cb["message"]["message_id"]
                    cb_data = cb.get("data", "")

                    if chat_type != "private":
                        answer_callback_query(
                            TELEGRAM_BOT_TOKEN, cb_id, 
                            "⚠️ Fitur ini hanya dapat digunakan di Chat Pribadi Bot!", 
                            show_alert=True
                        )
                        continue

                    if cb_data == "show_countries":
                        text, markup = get_countries_view()
                        if markup:
                            edit_telegram_message(TELEGRAM_BOT_TOKEN, chat_id, msg_id, text, reply_markup=markup)
                        else:
                            answer_callback_query(TELEGRAM_BOT_TOKEN, cb_id, text, show_alert=True)

                    elif cb_data.startswith("getnum_"):
                        country = cb_data.split("getnum_")[1]
                        text, markup = get_numbers_view(country)
                        if markup:
                            edit_telegram_message(TELEGRAM_BOT_TOKEN, chat_id, msg_id, text, reply_markup=markup)
                        else:
                            answer_callback_query(TELEGRAM_BOT_TOKEN, cb_id, text, show_alert=True)

                    continue

                # 2. MESSAGE HANDLER
                if "message" in update:
                    message = update["message"]
                    chat_id = message["chat"]["id"]
                    chat_type = message["chat"].get("type", "")
                    from_user = message.get("from", {})
                    user_id = from_user.get("id")
                    text = message.get("text", "").strip()
                    caption = message.get("caption", "").strip()

                    # FITUR UMUM DIIZINKAN DI GRUP ATAU PM
                    if text.startswith("/toprcv") or text.startswith("/top") or text == "📊 Top RCV":
                        send_top_rcv_report(target_chat=chat_id, reset=False)
                        continue

                    # BLOKIR INTERAKSI LAINNYA DI GRUP
                    if chat_type != "private":
                        if text in ["/start", "💥 Get Number", "🔥 Get File", "🏫 Chat Support"] or text.startswith("/"):
                            send_tg_reply(
                                TELEGRAM_BOT_TOKEN, chat_id, 
                                "⚠️ <b>Fitur ini hanya dapat digunakan di Chat Pribadi Bot!</b>\n"
                                "Silakan kirim pesan langsung ke Bot untuk menggunakan menu."
                            )
                        continue

                    # KHUSUS PRIVATE CHAT
                    if text == "/start":
                        welcome_text = "👑 <b>WELCOME TO OTP BOT</b>\n\nSilakan pilih menu di bawah ini:"
                        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, welcome_text, reply_markup=get_main_reply_keyboard())
                        continue

                    if text == "💥 Get Number":
                        txt, markup = get_whatsapp_stock_view()
                        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, txt, reply_markup=markup)
                        continue

                    if text == "🔥 Get File":
                        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "📁 Fitur Get File akan segera hadir!")
                        continue

                    if text == "🏫 Chat Support":
                        send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "💬 Silakan hubungi Administrator untuk bantuan.")
                        continue

                    # COMMAND KHUSUS ADMIN (PM ONLY)
                    if user_id in ADMIN_IDS:
                        if text.startswith("/add"):
                            parts = text.split()
                            if len(parts) < 3:
                                send_tg_reply(
                                    TELEGRAM_BOT_TOKEN, chat_id,
                                    "⚠️ <b>Format salah!</b>\n"
                                    "Gunakan: <code>/add <NEGARA> <NOMOR1> <NOMOR2> ...</code>\n\n"
                                    "Atau upload file <code>.txt</code> dengan caption <code>/add <NEGARA></code>."
                                )
                                continue

                            country_name = parts[1].upper()
                            numbers_to_add = parts[2:]

                            conn = sqlite3.connect(DB_PATH)
                            cursor = conn.cursor()
                            added, failed = 0, 0
                            for num in numbers_to_add:
                                try:
                                    cursor.execute("INSERT INTO numbers (country, phone_number) VALUES (?, ?)", (country_name, num))
                                    added += 1
                                except sqlite3.IntegrityError:
                                    failed += 1
                            conn.commit()
                            conn.close()

                            send_tg_reply(
                                TELEGRAM_BOT_TOKEN, chat_id,
                                f"✅ <b>PROSES TAMBAH NOMOR SELESAI</b>\n\n"
                                f"Negara: <b>{country_name}</b>\n"
                                f"Berhasil: <code>{added}</code>\n"
                                f"Gagal/Duplikat: <code>{failed}</code>"
                            )
                            continue

                        if text.startswith("/del") or text.startswith("/dell"):
                            parts = text.split(maxsplit=1)
                            if len(parts) < 2:
                                send_tg_reply(
                                    TELEGRAM_BOT_TOKEN, chat_id,
                                    "⚠️ <b>Format salah!</b>\n\n"
                                    "• Hapus 1 nomor: <code>/del +62812345678</code>\n"
                                    "• Hapus semua nomor 1 negara: <code>/dell INDONESIA</code>"
                                )
                                continue

                            target = parts[1].strip()
                            conn = sqlite3.connect(DB_PATH)
                            cursor = conn.cursor()

                            if target.startswith('+') or target.isdigit():
                                cursor.execute("DELETE FROM numbers WHERE phone_number = ?", (target,))
                                deleted_count = cursor.rowcount
                                conn.commit()
                                conn.close()
                                if deleted_count > 0:
                                    send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, f"✅ Nomor <code>{target}</code> berhasil dihapus.")
                                else:
                                    send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, f"❌ Nomor <code>{target}</code> tidak ditemukan.")
                            else:
                                country_name = target.upper()
                                cursor.execute("DELETE FROM numbers WHERE UPPER(country) = ?", (country_name,))
                                deleted_count = cursor.rowcount
                                conn.commit()
                                conn.close()
                                if deleted_count > 0:
                                    send_tg_reply(
                                        TELEGRAM_BOT_TOKEN, chat_id,
                                        f"🗑 <b>HAPUS MASSAL BERHASIL</b>\n\n"
                                        f"Negara: <b>{country_name}</b>\n"
                                        f"Total dihapus: <code>{deleted_count} nomor</code>"
                                    )
                                else:
                                    send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, f"❌ Tidak ditemukan stok nomor untuk negara <b>{country_name}</b>.")
                            continue

                        if text.startswith("/setchannel"):
                            parts = text.split(maxsplit=1)
                            if len(parts) < 2:
                                send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "⚠️ <b>Format salah!</b>\nGunakan: <code>/setchannel <LINK_CHANNEL></code>")
                                continue

                            new_link = parts[1].strip()
                            set_setting("otp_channel", new_link)
                            send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, f"✅ Link Channel OTP berhasil diperbarui ke:\n{new_link}")
                            continue

                        if "document" in message:
                            doc = message["document"]
                            file_id = doc["file_id"]
                            file_name = doc.get("file_name", "file").lower()

                            if file_name.endswith(".xlsx"):
                                handle_file_conversion(chat_id, file_id, file_name)
                                continue

                            if file_name.endswith(".txt") and caption.startswith("/add"):
                                parts = caption.split()
                                if len(parts) < 2:
                                    send_tg_reply(TELEGRAM_BOT_TOKEN, chat_id, "⚠️ Berikan caption: <code>/add <NEGARA></code> saat mengunggah file .txt")
                                    continue

                                country_name = parts[1].upper()

                                get_file_url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/getFile?file_id={file_id}"
                                res = http_session.get(get_file_url, timeout=10).json()
                                if res.get("ok"):
                                    file_path_tg = res["result"]["file_path"]
                                    download_url = f"https://api.telegram.org/file/bot{TELEGRAM_BOT_TOKEN}/{file_path_tg}"
                                    txt_content = http_session.get(download_url, timeout=30).text

                                    numbers = [line.strip() for line in txt_content.splitlines() if line.strip()]

                                    conn = sqlite3.connect(DB_PATH)
                                    cursor = conn.cursor()
                                    added, failed = 0, 0
                                    for num in numbers:
                                        try:
                                            cursor.execute("INSERT INTO numbers (country, phone_number) VALUES (?, ?)", (country_name, num))
                                            added += 1
                                        except sqlite3.IntegrityError:
                                            failed += 1
                                    conn.commit()
                                    conn.close()

                                    send_tg_reply(
                                        TELEGRAM_BOT_TOKEN, chat_id,
                                        f"✅ <b>IMPORT FILE TXT SELESAI</b>\n\n"
                                        f"Negara: <b>{country_name}</b>\n"
                                        f"Berhasil: <code>{added}</code>\n"
                                        f"Gagal/Duplikat: <code>{failed}</code>"
                                    )
                                continue

        except Exception as e:
            print("⚠️ Error pada Telegram Listener:", e)
            time.sleep(2)

# ==============================================================================
# MAIN THREADING LOOP & ENTRY POINT
# ==============================================================================
def loop_fetch_api_1():
    while True:
        fetch_api_1()
        time.sleep(CHECK_INTERVAL)

def loop_fetch_api_2():
    while True:
        fetch_api_2()
        time.sleep(KSIIPRN_CHECK_INTERVAL)

def loop_fetch_api_3():
    while True:
        fetch_api_3()
        time.sleep(AUGESTEL_CHECK_INTERVAL)

def loop_fetch_api_4():
    while True:
        fetch_api_4()
        time.sleep(THIRDWAVE_CHECK_INTERVAL)

def loop_top_rcv_report():
    global last_top_rcv_time
    while True:
        if time.time() - last_top_rcv_time >= TOP_RCV_INTERVAL:
            send_top_rcv_report(target_chat=CHAT_ID, reset=True)
            last_top_rcv_time = time.time()
        time.sleep(60)

if __name__ == "__main__":
    print("🚀 Memulai Bot & Service Monitoring OTP...")
    delete_webhook(TELEGRAM_BOT_TOKEN)

    threading.Thread(target=loop_fetch_api_1, daemon=True).start()
    threading.Thread(target=loop_fetch_api_2, daemon=True).start()
    threading.Thread(target=loop_fetch_api_3, daemon=True).start()
    threading.Thread(target=loop_fetch_api_4, daemon=True).start()
    threading.Thread(target=loop_top_rcv_report, daemon=True).start()

    telegram_listener()
